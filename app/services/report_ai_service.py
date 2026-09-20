"""
Servicio de Asistente de Inteligencia Artificial para Reportes Ejecutivos (StyleStore BI).
Extrae datos operativos en tiempo real de la base de datos (ventas del día, recaudación
por sucursal, pedidos por delivery, prendas vendidas, inventario crítico y formas de pago)
y consulta el motor Groq LLM para responder analíticamente a preguntas del administrador/encargado.
"""
import re
import json
from datetime import date, datetime, timedelta
from typing import Dict, Any, List, Optional
from decimal import Decimal
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func, or_

from app.models.orden_venta import OrdenVenta, DetalleVenta
from app.models.sucursal import Sucursal
from app.models.producto import Producto
from app.models.producto_color import ProductoColor
from app.models.envio import Envio
from app.models.stock_inventario import StockInventario
from app.models.cliente import Cliente
from app.config import settings


class ReportAIService:
    """Motor analítico de IA para consultas gerenciales y ejecutivas en Reportes."""

    def __init__(self, db: Session):
        self.db = db

    def _recopilar_contexto_operativo(self) -> Dict[str, Any]:
        """Extrae un resumen consolidado de las operaciones de StyleStore para el LLM."""
        hoy = date.today()
        hace_7_dias = hoy - timedelta(days=7)
        primer_dia_mes = hoy.replace(day=1)

        # 1. Ventas de HOY
        ordenes_hoy = []
        try:
            ordenes_hoy = (
                self.db.query(OrdenVenta)
                .options(joinedload(OrdenVenta.sucursal), joinedload(OrdenVenta.detalles))
                .filter(OrdenVenta.fecha == hoy, OrdenVenta.estado != "cancelada")
                .all()
            )
        except Exception:
            pass

        total_hoy = sum([float(o.total) for o in ordenes_hoy]) if ordenes_hoy else 0.0
        prendas_vendidas_detalle = []
        total_prendas_hoy = 0

        for o in ordenes_hoy:
            if not o.detalles:
                continue
            for d in o.detalles:
                cant = int(d.cantidad) if d.cantidad else 1
                total_prendas_hoy += cant
                nom = getattr(d, "producto_nombre", None) or "Prenda StyleStore"
                color = getattr(d, "color_nombre", None)
                talla = getattr(d, "talla_nombre", None)
                sub = float(d.subtotal) if getattr(d, "subtotal", None) else 0.0

                extra = []
                if color:
                    extra.append(f"Color: {color}")
                if talla:
                    extra.append(f"Talla: {talla}")
                extra_str = f" ({', '.join(extra)})" if extra else ""

                prendas_vendidas_detalle.append(f"{cant}x {nom}{extra_str} - Bs. {sub:.2f}")

        # 2. Recaudación por Sucursal (Hoy y Mes)
        sucursales = self.db.query(Sucursal).filter(Sucursal.active == True).all()
        sucursal_metricas = []
        for s in sucursales:
            recaudacion_hoy_suc = sum(
                [float(o.total) for o in ordenes_hoy if o.sucursal_id == s.id]
            )
            # Recaudación en el mes
            ordenes_mes_suc = (
                self.db.query(func.coalesce(func.sum(OrdenVenta.total), 0))
                .filter(
                    OrdenVenta.sucursal_id == s.id,
                    OrdenVenta.fecha >= primer_dia_mes,
                    OrdenVenta.estado != "cancelada",
                )
                .scalar()
            )
            sucursal_metricas.append({
                "id": s.id,
                "nombre": s.nombre,
                "ciudad": s.ciudad,
                "direccion": s.direccion,
                "recaudado_hoy": float(recaudacion_hoy_suc),
                "recaudado_mes": float(ordenes_mes_suc or 0.0),
            })

        # 3. Pedidos por Delivery / Envíos
        envios_hoy = self.db.query(Envio).filter(Envio.fecha == hoy).all()
        total_envios_hoy = len(envios_hoy)
        entregados_hoy = len([e for e in envios_hoy if e.estado in ["entregado", "completado"]])
        en_camino_hoy = len([e for e in envios_hoy if e.estado in ["en camino", "en_camino"]])
        pendientes_hoy = len([e for e in envios_hoy if e.estado in ["pendiente"]])
        total_tarifas_delivery_hoy = sum([float(e.costo) for e in envios_hoy])

        envios_mes_total = (
            self.db.query(func.count(Envio.id))
            .filter(Envio.fecha >= primer_dia_mes)
            .scalar()
            or 0
        )

        # 4. Inventario Crítico (Prendas con stock <= 5 o agotadas)
        stocks_bajos = (
            self.db.query(StockInventario)
            .options(
                joinedload(StockInventario.producto_color).joinedload(ProductoColor.producto),
                joinedload(StockInventario.sucursal),
                joinedload(StockInventario.talla),
            )
            .filter(StockInventario.cantidad <= 5)
            .limit(20)
            .all()
        )
        inventario_critico_list = []
        for st in stocks_bajos:
            prod = st.producto_color.producto if (st.producto_color and st.producto_color.producto) else None
            p_nom = prod.nombre if prod else "Prenda"
            p_cod = prod.codigo if prod else "N/A"
            suc_nom = st.sucursal.nombre if st.sucursal else "General"
            talla_nom = st.talla.nombre if st.talla else "Única"
            inventario_critico_list.append(
                f"{p_nom} (Ref: {p_cod}, Talla: {talla_nom}, Sucursal: {suc_nom}) -> Stock restante: {st.cantidad}"
            )

        # 5. Métodos de Pago Hoy
        pagos_hoy_map = {}
        for o in ordenes_hoy:
            met = (o.metodo_pago or "EFECTIVO").upper()
            pagos_hoy_map[met] = pagos_hoy_map.get(met, 0.0) + float(o.total)

        # 6. Recaudación Total Mes
        recaudacion_mes_total = (
            self.db.query(func.coalesce(func.sum(OrdenVenta.total), 0))
            .filter(OrdenVenta.fecha >= primer_dia_mes, OrdenVenta.estado != "cancelada")
            .scalar()
        )

        return {
            "fecha_actual": hoy.isoformat(),
            "ventas_hoy": {
                "total_bs": total_hoy,
                "cantidad_ordenes": len(ordenes_hoy),
                "total_prendas": total_prendas_hoy,
                "prendas_detalle": prendas_vendidas_detalle,
            },
            "sucursales": sucursal_metricas,
            "delivery": {
                "total_hoy": total_envios_hoy,
                "entregados_hoy": entregados_hoy,
                "en_camino_hoy": en_camino_hoy,
                "pendientes_hoy": pendientes_hoy,
                "recaudado_tarifas_hoy": total_tarifas_delivery_hoy,
                "total_mes": envios_mes_total,
            },
            "inventario_critico": inventario_critico_list,
            "pagos_hoy": pagos_hoy_map,
            "recaudacion_mes_total": float(recaudacion_mes_total or 0.0),
        }

    def procesar_consulta(self, pregunta: str, user_name: str = "Administrador") -> Dict[str, Any]:
        """Procesa la pregunta analítica mediante Groq LLM o motor de cálculo local."""
        pregunta_clean = pregunta.strip()
        contexto = self._recopilar_contexto_operativo()

        if settings.GROQ_API_KEY:
            try:
                import httpx

                # Convertir contexto a texto conciso para el prompt
                suc_txt = "\n".join([
                    f"- {s['nombre']} ({s['ciudad']}): Hoy Bs. {s['recaudado_hoy']:.2f} | Mes Bs. {s['recaudado_mes']:.2f}"
                    for s in contexto["sucursales"]
                ]) or "Sin sucursales registradas"

                prendas_txt = (
                    "\n".join([f"- {p}" for p in contexto["ventas_hoy"]["prendas_detalle"]])
                    if contexto["ventas_hoy"]["prendas_detalle"]
                    else "No se han registrado ventas de prendas en el día de hoy hasta el momento."
                )

                critico_txt = (
                    "\n".join([f"- {item}" for item in contexto["inventario_critico"][:8]])
                    if contexto["inventario_critico"]
                    else "Todos los niveles de inventario se encuentran estables (por encima de 5 unidades)."
                )

                pagos_txt = ", ".join([f"{k}: Bs. {v:.2f}" for k, v in contexto["pagos_hoy"].items()]) or "Sin registros hoy"

                system_prompt = f"""Eres el Asistente Ejecutivo de Inteligencia Artificial para el Módulo de Reportes de "StyleStore" (Boutique de moda en Bolivia).
Tu interlocutor es el usuario administrador o encargado de tienda: {user_name}.
Tu misión es responder con absoluta precisión analítica, ejecutiva, elegante y directa a preguntas sobre ventas, sucursales, pedidos por delivery, inventario y métricas clave.
La moneda oficial es el Boliviano (Bs.). La fecha actual del sistema es {contexto['fecha_actual']}.

SNAPSHOT DE DATOS OPERATIVOS EN VIVO:
1. VENTAS DE HOY:
- Recaudación total hoy: Bs. {contexto['ventas_hoy']['total_bs']:.2f}
- Cantidad de órdenes hoy: {contexto['ventas_hoy']['cantidad_ordenes']}
- Total de prendas vendidas hoy: {contexto['ventas_hoy']['total_prendas']}
- Lista de prendas vendidas hoy:
{prendas_txt}

2. DESGLOSE POR SUCURSAL:
{suc_txt}

3. PEDIDOS POR DELIVERY Y ENVÍOS:
- Total pedidos delivery hoy: {contexto['delivery']['total_hoy']} (Entregados: {contexto['delivery']['entregados_hoy']}, En camino: {contexto['delivery']['en_camino_hoy']}, Pendientes: {contexto['delivery']['pendientes_hoy']})
- Total cobrado en tarifas de delivery hoy: Bs. {contexto['delivery']['recaudado_tarifas_hoy']:.2f}
- Pedidos delivery acumulados en el mes: {contexto['delivery']['total_mes']}

4. MÉTODOS DE PAGO HOY:
{pagos_txt}

5. RECAUDACIÓN ACUMULADA DEL MES:
- Total general del mes: Bs. {contexto['recaudacion_mes_total']:.2f}

6. INVENTARIO CRÍTICO (STOCK <= 5 O AGOTADOS):
{critico_txt}

DIRECTRICES DE RESPUESTA:
- Responde siempre en español, con tono formal, ejecutivo y claro.
- Proporciona cifras exactas. Si preguntan "¿cuántas ropas se vendieron hoy? y cuáles son?", menciona el total de prendas y la lista de prendas exactas vendidas hoy.
- Si preguntan por una sucursal específica (ej. "¿cuánto recaudó la sucursal alemana?"), busca la coincidencia por nombre o ciudad y reporta tanto lo recaudado hoy como en el mes.
- Si preguntan por delivery, detalla la cantidad de pedidos y los estados.
- Si no hay datos registrados hoy (ej. 0 ventas o sin órdenes hoy), explícalo con claridad indicando que no se registran movimientos en la fecha actual aún.
- Al final de tu respuesta, si aplica, incluye una línea especial con 1 a 3 tarjetas métricas KPI en formato JSON:
KPIS: [{{"label": "Vendido Hoy", "valor": "Bs. 0.00"}}, {{"label": "Prendas Vendidas", "valor": "0"}}]
"""

                response = httpx.post(
                    "https://api.groq.com/openai/v1/chat/completions",
                    headers={
                        "Authorization": f"Bearer {settings.GROQ_API_KEY}",
                        "Content-Type": "application/json",
                    },
                    json={
                        "model": settings.GROQ_MODEL,
                        "messages": [
                            {"role": "system", "content": system_prompt},
                            {"role": "user", "content": pregunta_clean},
                        ],
                        "temperature": 0.2,
                        "max_tokens": 600,
                    },
                    timeout=12.0,
                )

                if response.status_code == 200:
                    data = response.json()
                    content = data.get("choices", [{}])[0].get("message", {}).get("content", "").strip()

                    kpis = []
                    text_lines = []
                    skip_code_block = False
                    for line in content.splitlines():
                        s = line.strip()
                        if s.startswith("KPIS:"):
                            kpi_part = s[5:].strip()
                            try:
                                parsed_kpis = json.loads(kpi_part)
                                if isinstance(parsed_kpis, list):
                                    kpis = parsed_kpis
                            except Exception:
                                pass
                            continue
                        if s.lower().startswith("**kpis"):
                            continue
                        if s.startswith("```json"):
                            skip_code_block = True
                            continue
                        if skip_code_block and s == "```":
                            skip_code_block = False
                            continue
                        if not skip_code_block:
                            text_lines.append(line)

                    cleaned_text = "\n".join(text_lines).strip()

                    if not kpis:
                        # KPIs automáticos según contexto
                        kpis = [
                            {"label": "Total Hoy", "valor": f"Bs. {contexto['ventas_hoy']['total_bs']:.2f}"},
                            {"label": "Prendas Hoy", "valor": str(contexto['ventas_hoy']['total_prendas'])},
                            {"label": "Delivery Hoy", "valor": str(contexto['delivery']['total_hoy'])},
                        ]

                    return {
                        "pregunta": pregunta_clean,
                        "respuesta": cleaned_text,
                        "kpis": kpis,
                        "ia_powered": True,
                    }
                else:
                    print(f"[ReportAIService] Groq API returned status {response.status_code}: {response.text}")
            except Exception as e:
                print(f"[ReportAIService] Exception contacting Groq LLM: {e}")
                pass

        # Fallback Local en caso de que Groq esté inactivo
        return self._generar_respuesta_local(pregunta_clean, contexto)

    def _generar_respuesta_local(self, pregunta: str, ctx: Dict[str, Any]) -> Dict[str, Any]:
        """Genera respuesta determinística local si el LLM no responde."""
        p_lower = pregunta.lower()

        # Prioridad 1: Delivery y Envíos
        if any(w in p_lower for w in ["delivery", "envio", "envíos", "envios", "repartidor", "moto", "despacho"]):
            d = ctx["delivery"]
            lineas = [
                f"🚚 **Reporte de Envíos y Delivery:**",
                f"• **Pedidos por Delivery Hoy:** {d['total_hoy']}",
                f"• **Entregados con Éxito:** {d['entregados_hoy']}",
                f"• **En Camino (Rastreo en Vivo):** {d['en_camino_hoy']}",
                f"• **Pendientes de Asignación:** {d['pendientes_hoy']}",
                f"• **Tarifas Recaudadas Hoy:** Bs. {d['recaudado_tarifas_hoy']:.2f}",
                f"• **Total Envíos en el Mes:** {d['total_mes']}",
            ]
            return {
                "pregunta": pregunta,
                "respuesta": "\n".join(lineas),
                "kpis": [
                    {"label": "Delivery Hoy", "valor": str(d["total_hoy"])},
                    {"label": "Entregados", "valor": str(d["entregados_hoy"])},
                    {"label": "En Camino", "valor": str(d["en_camino_hoy"])},
                ],
                "ia_powered": False,
            }

        # Prioridad 2: Sucursales específicas o listado de sucursales
        for s in ctx["sucursales"]:
            if s["nombre"].lower() in p_lower or (len(s["ciudad"]) > 3 and s["ciudad"].lower() in p_lower):
                return {
                    "pregunta": pregunta,
                    "respuesta": f"🏢 **Reporte Financiero de {s['nombre']} ({s['ciudad']}):**\n"
                                 f"• **Dirección:** {s['direccion']}\n"
                                 f"• **Recaudación Hoy:** Bs. {s['recaudado_hoy']:.2f}\n"
                                 f"• **Recaudación Acumulada del Mes:** Bs. {s['recaudado_mes']:.2f}",
                    "kpis": [
                        {"label": f"Hoy {s['nombre']}", "valor": f"Bs. {s['recaudado_hoy']:.2f}"},
                        {"label": f"Mes {s['nombre']}", "valor": f"Bs. {s['recaudado_mes']:.2f}"},
                    ],
                    "ia_powered": False,
                }

        if any(w in p_lower for w in ["sucursal", "sucursales", "tienda", "tiendas"]):
            lineas = ["🏢 **Desglose de Recaudación por Sucursales:**\n"]
            for s in ctx["sucursales"]:
                lineas.append(f"• **{s['nombre']} ({s['ciudad']}):** Hoy Bs. {s['recaudado_hoy']:.2f} | Mes Bs. {s['recaudado_mes']:.2f}")
            return {
                "pregunta": pregunta,
                "respuesta": "\n".join(lineas),
                "kpis": [{"label": "Sucursales Activas", "valor": str(len(ctx["sucursales"]))}],
                "ia_powered": False,
            }

        # Prioridad 3: Stock e Inventario Crítico
        if any(w in p_lower for w in ["stock", "inventario", "agotado", "agotados", "quedan", "faltante", "crítico", "critico"]):
            crit = ctx["inventario_critico"]
            if crit:
                lineas = ["⚠️ **Prendas con Stock Crítico (5 unidades o menos):**\n"]
                for c in crit[:12]:
                    lineas.append(f"• {c}")
                return {
                    "pregunta": pregunta,
                    "respuesta": "\n".join(lineas),
                    "kpis": [{"label": "Prendas Críticas", "valor": str(len(crit))}],
                    "ia_powered": False,
                }
            else:
                return {
                    "pregunta": pregunta,
                    "respuesta": "✅ **Estado Óptimo de Inventario:**\nTodas las prendas registradas en las sucursales cuentan con stock suficiente (más de 5 unidades disponibles).",
                    "kpis": [{"label": "Estado Stock", "valor": "Óptimo"}],
                    "ia_powered": False,
                }

        # Prioridad 4: Métodos de Pago y Finanzas
        if any(w in p_lower for w in ["pago", "pagos", "metodo", "método", "qr", "efectivo", "tarjeta", "paypal", "cobro", "cobrado"]):
            pagos = ctx["pagos_hoy"]
            lineas = ["💳 **Distribución de Métodos de Pago de Hoy:**\n"]
            if pagos:
                for met, monto in pagos.items():
                    lineas.append(f"• **{met.upper()}:** Bs. {monto:.2f}")
            else:
                lineas.append("No se registran transacciones cobradas en el día de hoy hasta el momento.")
            return {
                "pregunta": pregunta,
                "respuesta": "\n".join(lineas),
                "kpis": [
                    {"label": "Canales Activos", "valor": str(len(pagos))},
                    {"label": "Total Hoy", "valor": f"Bs. {ctx['ventas_hoy']['total_bs']:.2f}"},
                ],
                "ia_powered": False,
            }

        # Prioridad 5: Recaudación del Mes / Acumulado
        if any(w in p_lower for w in ["mes", "mensual", "acumulado"]):
            return {
                "pregunta": pregunta,
                "respuesta": f"📈 **Recaudación Acumulada del Mes:**\n"
                             f"• **Total Facturado en el Mes:** Bs. {ctx['recaudacion_mes_total']:.2f}\n"
                             f"• **Ventas Registradas Hoy:** Bs. {ctx['ventas_hoy']['total_bs']:.2f} ({ctx['ventas_hoy']['total_prendas']} prendas)",
                "kpis": [
                    {"label": "Total Mes", "valor": f"Bs. {ctx['recaudacion_mes_total']:.2f}"},
                    {"label": "Total Hoy", "valor": f"Bs. {ctx['ventas_hoy']['total_bs']:.2f}"},
                ],
                "ia_powered": False,
            }

        # Prioridad 6: Ventas de Hoy / Prendas Vendidas Hoy
        if any(w in p_lower for w in ["prenda", "prendas", "ropa", "ropas", "vendieron", "vendido", "orden", "ordenes", "órdenes", "recaudo", "recaudó", "recaudación", "recaudacion", "ventas", "venta", "hoy"]):
            total_bs = ctx["ventas_hoy"]["total_bs"]
            total_prendas = ctx["ventas_hoy"]["total_prendas"]
            prendas = ctx["ventas_hoy"]["prendas_detalle"]

            lineas = [
                f"📊 **Reporte Operativo de Ventas de Hoy:**",
                f"• **Recaudación Total:** Bs. {total_bs:.2f}",
                f"• **Cantidad de Órdenes:** {ctx['ventas_hoy']['cantidad_ordenes']}",
                f"• **Prendas Vendidas:** {total_prendas} unidades\n",
            ]
            if prendas:
                lineas.append("**Detalle de prendas vendidas:**")
                for pr in prendas:
                    lineas.append(f"• {pr}")
            else:
                lineas.append("Hasta el momento no se registran prendas vendidas en el sistema el día de hoy.")

            return {
                "pregunta": pregunta,
                "respuesta": "\n".join(lineas),
                "kpis": [
                    {"label": "Recaudado Hoy", "valor": f"Bs. {total_bs:.2f}"},
                    {"label": "Prendas", "valor": str(total_prendas)},
                ],
                "ia_powered": False,
            }

        # Prioridad 7: Respuesta general consolidada
        return {
            "pregunta": pregunta,
            "respuesta": f"📈 **Resumen Ejecutivo StyleStore:**\n"
                         f"• Recaudación Hoy: Bs. {ctx['ventas_hoy']['total_bs']:.2f} ({ctx['ventas_hoy']['total_prendas']} prendas)\n"
                         f"• Recaudación Acumulada Mes: Bs. {ctx['recaudacion_mes_total']:.2f}\n"
                         f"• Pedidos Delivery Hoy: {ctx['delivery']['total_hoy']}\n\n"
                         f"Puedes consultar por prendas vendidas hoy, recaudación por sucursal, pedidos delivery, formas de pago o stock crítico.",
            "kpis": [
                {"label": "Ventas Hoy", "valor": f"Bs. {ctx['ventas_hoy']['total_bs']:.2f}"},
                {"label": "Mes Total", "valor": f"Bs. {ctx['recaudacion_mes_total']:.2f}"},
            ],
            "ia_powered": False,
        }

    def exportar_excel(self, pregunta: str, respuesta: str, kpis: List[Dict[str, Any]]) -> Any:
        """Genera un archivo Excel (.xlsx) con la consulta ejecutiva formulada a la IA."""
        import io
        import openpyxl  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]

        wb = openpyxl.Workbook()
        ws: Any = wb.active if wb.active is not None else wb.create_sheet()
        assert ws is not None
        ws.title = "Consulta Ejecutiva IA"

        navy_fill = PatternFill(start_color="14263D", end_color="14263D", fill_type="solid")
        gold_fill = PatternFill(start_color="C6A87D", end_color="C6A87D", fill_type="solid")
        light_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        title_font = Font(name="Arial", size=14, bold=True, color="14263D")
        sub_font = Font(name="Arial", size=9, italic=True, color="64748B")
        section_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        normal_font = Font(name="Arial", size=10)
        bold_font = Font(name="Arial", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0'),
        )

        ws.merge_cells("A1:D1")
        ws["A1"] = "StyleStore - Reporte Ejecutivo Asistente IA"
        ws["A1"].font = title_font

        ws.merge_cells("A2:D2")
        ws["A2"] = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Módulo de Inteligencia Artificial & BI"
        ws["A2"].font = sub_font

        # Pregunta
        ws.merge_cells("A4:D4")
        ws["A4"] = "CONSULTA FORMULADA:"
        ws["A4"].fill = navy_fill
        ws["A4"].font = section_font
        ws["A4"].alignment = Alignment(vertical="center", indent=1)

        ws.merge_cells("A5:D5")
        ws["A5"] = f'"{pregunta}"'
        ws["A5"].font = bold_font
        ws["A5"].alignment = Alignment(vertical="center", indent=1)
        ws["A5"].fill = light_fill

        curr_row = 7
        # KPIs
        if kpis:
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=2)
            ws.cell(row=curr_row, column=1, value="MÉTRICAS CLAVE (KPIS)").fill = gold_fill
            ws.cell(row=curr_row, column=1).font = Font(name="Arial", size=10, bold=True, color="14263D")
            curr_row += 1

            ws.cell(row=curr_row, column=1, value="Indicador").fill = navy_fill
            ws.cell(row=curr_row, column=1).font = header_font
            ws.cell(row=curr_row, column=2, value="Valor").fill = navy_fill
            ws.cell(row=curr_row, column=2).font = header_font
            curr_row += 1

            for k in kpis:
                c1 = ws.cell(row=curr_row, column=1, value=k.get("label", ""))
                c2 = ws.cell(row=curr_row, column=2, value=k.get("valor", ""))
                c1.font = normal_font
                c2.font = bold_font
                c1.border = thin_border
                c2.border = thin_border
                curr_row += 1
            curr_row += 1

        # Respuesta Analítica
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
        ws.cell(row=curr_row, column=1, value="INFORME ANALÍTICO DE LA CONSULTA").fill = navy_fill
        ws.cell(row=curr_row, column=1).font = section_font
        ws.cell(row=curr_row, column=1).alignment = Alignment(vertical="center", indent=1)
        curr_row += 1

        for line in respuesta.splitlines():
            line_clean = line.strip()
            if not line_clean:
                curr_row += 1
                continue
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
            c = ws.cell(row=curr_row, column=1, value=line_clean)
            c.font = bold_font if line_clean.startswith("•") or line_clean.startswith("📊") or line_clean.startswith("🚚") or line_clean.startswith("🏢") or line_clean.startswith("⚠️") or line_clean.startswith("💳") else normal_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
            curr_row += 1

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 25

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def exportar_pdf(self, pregunta: str, respuesta: str, kpis: List[Dict[str, Any]]) -> Any:
        """Genera un archivo PDF (.pdf) estilizado con el reporte analítico de la IA."""
        import io
        import re
        from reportlab.lib.pagesizes import letter  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
        from reportlab.lib import colors  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'RepAITitle',
            parent=styles['Heading1'],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#14263D'),
            fontName="Helvetica-Bold",
        )
        subtitle_style = ParagraphStyle(
            'RepAISub',
            parent=styles['Normal'],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#64748B'),
        )
        query_label_style = ParagraphStyle(
            'RepAIQueryLbl',
            parent=styles['Normal'],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#94A3B8'),
            fontName="Helvetica-Bold",
        )
        query_text_style = ParagraphStyle(
            'RepAIQueryText',
            parent=styles['Normal'],
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#14263D'),
            fontName="Helvetica-Bold",
        )
        heading2_style = ParagraphStyle(
            'RepAIHead2',
            parent=styles['Heading2'],
            fontSize=12,
            leading=15,
            textColor=colors.HexColor('#14263D'),
            fontName="Helvetica-Bold",
        )
        body_style = ParagraphStyle(
            'RepAIBody',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#1E293B'),
        )
        bullet_style = ParagraphStyle(
            'RepAIBullet',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#0F172A'),
            leftIndent=12,
        )

        elements: List[Any] = [
            Paragraph("StyleStore - Reporte Ejecutivo Asistente IA", title_style),
            Paragraph(f"Emitido el: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Sistema StyleStore BI", subtitle_style),
            Spacer(1, 14),
        ]

        # Caja de la pregunta
        query_table = Table(
            [
                [Paragraph("CONSULTA EJECUTIVA FORMULADA:", query_label_style)],
                [Paragraph(f'"{pregunta}"', query_text_style)]
            ],
            colWidths=[540],
        )
        query_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#E2E8F0')),
            ('LINELEFT', (0, 0), (0, -1), 4, colors.HexColor('#14263D')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ]))
        elements.append(query_table)
        elements.append(Spacer(1, 14))

        # KPIs si existen
        if kpis:
            elements.append(Paragraph("MÉTRICAS CLAVE (KPIS)", heading2_style))
            elements.append(Spacer(1, 6))

            kpi_table_data = [["Indicador / Métrica", "Valor Reportado"]]
            for k in kpis:
                kpi_table_data.append([k.get("label", "-"), k.get("valor", "-")])

            kpi_table = Table(kpi_table_data, colWidths=[270, 270])
            kpi_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#14263D')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F8F9FA'), colors.white]),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ]))
            elements.append(kpi_table)
            elements.append(Spacer(1, 14))

        # Respuesta Analítica
        elements.append(Paragraph("INFORME ANALÍTICO DETALLADO", heading2_style))
        elements.append(Spacer(1, 6))

        for line in respuesta.splitlines():
            line_str = line.strip()
            if not line_str:
                elements.append(Spacer(1, 4))
                continue
            clean_str = line_str.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            clean_str = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', clean_str)

            if clean_str.startswith("•") or clean_str.startswith("-"):
                elements.append(Paragraph(clean_str, bullet_style))
            else:
                elements.append(Paragraph(clean_str, body_style))
            elements.append(Spacer(1, 3))

        doc.build(elements)
        buffer.seek(0)
        return buffer
