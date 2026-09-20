"""
Servicio de Reportes Dinámicos StyleStore.
Genera reportes de Ventas, Inventario y Clientes en Excel (.xlsx con openpyxl) y PDF (.pdf con reportlab).
"""
import io
import re
from datetime import datetime, date
from decimal import Decimal
from typing import List, Optional, Any, Dict
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func

import openpyxl  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]

from reportlab.lib.pagesizes import letter  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
from reportlab.lib import colors  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle  # pyrefly: ignore[untyped-import]  # type: ignore[import-untyped]

from app.models.orden_venta import OrdenVenta, DetalleVenta
from app.models.producto import Producto
from app.models.producto_color import ProductoColor
from app.models.color import Color
from app.models.talla import Talla
from app.models.stock_inventario import StockInventario
from app.models.sucursal import Sucursal
from app.models.cliente import Cliente
from app.models.user import User
from app.models.empleado import Empleado
from app.models.envio import Envio
from app.models.cambio_devolucion import CambioDevolucion
from app.models.categoria import Categoria
from app.models.pago import Pago
from app.models.temporada import Temporada
from app.models.bitacora import Bitacora


def clean_sheet_title(title: str) -> str:
    """Remueve caracteres no permitidos en nombres de hojas de Excel y trunca a 31 caracteres."""
    cleaned = re.sub(r'[\\/*?:\[\]]', '', title or "Reporte")
    return cleaned[:31]


def _get_cliente_nombre(cliente) -> str:
    """Extrae el nombre completo del cliente o su código de forma segura."""
    if not cliente:
        return "Cliente General"
    if hasattr(cliente, "user") and cliente.user:
        full = f"{cliente.user.name} {cliente.user.apellido or ''}".strip()
        if full:
            return full
    if hasattr(cliente, "nombre") and cliente.nombre:
        return cliente.nombre
    return getattr(cliente, "codigo", "Cliente General")


def _get_cliente_email(cliente) -> str:
    """Extrae el correo electrónico del cliente de forma segura."""
    if not cliente:
        return ""
    if hasattr(cliente, "user") and cliente.user and getattr(cliente.user, "email", None):
        return cliente.user.email
    if hasattr(cliente, "email") and cliente.email:
        return str(cliente.email)
    return ""


def _get_sucursal_nombre(sucursal) -> str:
    """Extrae el nombre o ciudad de la sucursal de forma segura."""
    if not sucursal:
        return "Central"
    if hasattr(sucursal, "name") and sucursal.name:
        return sucursal.name
    if hasattr(sucursal, "city") and sucursal.city:
        return sucursal.city
    return getattr(sucursal, "nombre", "Central")


class ReportService:
    """Generador de reportes en Excel y PDF con estética StyleStore."""

    def __init__(self, db: Session):
        self.db = db

    # ==========================================
    # 1. REPORTES DE VENTAS
    # ==========================================

    def get_ventas_data(
        self,
        fecha_inicio: Optional[date] = None,
        fecha_fin: Optional[date] = None,
        sucursal_id: Optional[int] = None,
        metodo_pago: Optional[str] = None,
    ) -> List[OrdenVenta]:
        """Obtiene las ventas filtradas."""
        query = self.db.query(OrdenVenta).options(
            joinedload(OrdenVenta.cliente),
            joinedload(OrdenVenta.sucursal),
            joinedload(OrdenVenta.detalles),
        )
        if fecha_inicio:
            query = query.filter(OrdenVenta.fecha >= fecha_inicio)
        if fecha_fin:
            query = query.filter(OrdenVenta.fecha <= fecha_fin)
        if sucursal_id:
            query = query.filter(OrdenVenta.sucursal_id == sucursal_id)
        if metodo_pago:
            query = query.filter(OrdenVenta.metodo_pago == metodo_pago)

        return query.order_by(OrdenVenta.fecha.desc(), OrdenVenta.id.desc()).all()

    def generate_ventas_excel(
        self,
        fecha_inicio: Optional[date] = None,
        fecha_fin: Optional[date] = None,
        sucursal_id: Optional[int] = None,
        metodo_pago: Optional[str] = None,
    ) -> io.BytesIO:
        """Genera un archivo Excel con el reporte de ventas."""
        ventas = self.get_ventas_data(fecha_inicio, fecha_fin, sucursal_id, metodo_pago)

        wb = openpyxl.Workbook()
        ws: Any = wb.active if wb.active is not None else wb.create_sheet()
        assert ws is not None
        ws.title = clean_sheet_title("Ventas StyleStore")

        # Paleta de colores
        navy_fill = PatternFill(start_color="14263D", end_color="14263D", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=16, bold=True, color="14263D")
        bold_font = Font(name="Arial", size=10, bold=True)
        normal_font = Font(name="Arial", size=10)
        thin_border = Border(
            left=Side(style='thin', color='E0E0E0'),
            right=Side(style='thin', color='E0E0E0'),
            top=Side(style='thin', color='E0E0E0'),
            bottom=Side(style='thin', color='E0E0E0')
        )

        # Encabezado del reporte
        ws.merge_cells("A1:H1")
        ws["A1"] = "StyleStore - Reporte Detallado de Ventas"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        gen_date = datetime.now().strftime("%Y-%m-%d %H:%M")
        ws["A2"] = f"Generado: {gen_date} | Total Registros: {len(ventas)}"
        ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")

        # Columnas
        headers = ["N° Venta / Ticket", "Fecha", "Cliente", "Correo Electrónico", "Sucursal", "Método Pago", "Estado", "Total (Bs.)"]
        ws.row_dimensions[4].height = 24

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Filas de datos
        row_idx = 5
        gran_total = Decimal("0.00")

        for v in ventas:
            total_num = float(v.total)
            gran_total += Decimal(str(v.total))
            cli_nom = _get_cliente_nombre(v.cliente)
            cli_email = _get_cliente_email(v.cliente)
            suc_nom = _get_sucursal_nombre(v.sucursal)

            ws.cell(row=row_idx, column=1, value=v.ticket_numero or f"ORD-{v.id:04d}").font = normal_font
            ws.cell(row=row_idx, column=2, value=v.fecha.strftime("%Y-%m-%d") if v.fecha else "").font = normal_font
            ws.cell(row=row_idx, column=3, value=cli_nom).font = normal_font
            ws.cell(row=row_idx, column=4, value=cli_email or "-").font = normal_font
            ws.cell(row=row_idx, column=5, value=suc_nom).font = normal_font
            ws.cell(row=row_idx, column=6, value=(v.metodo_pago or "EFECTIVO").upper()).font = normal_font
            ws.cell(row=row_idx, column=7, value=v.estado.capitalize()).font = normal_font
            
            c_total = ws.cell(row=row_idx, column=8, value=total_num)
            c_total.font = normal_font
            c_total.number_format = 'Bs. #,##0.00'
            c_total.alignment = Alignment(horizontal="right")

            for c in range(1, 9):
                ws.cell(row=row_idx, column=c).border = thin_border
            row_idx += 1

        # Fila de Total
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
        tot_label = ws.cell(row=row_idx, column=1, value="TOTAL GENERAL:")
        tot_label.font = bold_font
        tot_label.alignment = Alignment(horizontal="right", vertical="center")
        
        tot_val = ws.cell(row=row_idx, column=8, value=float(gran_total))
        tot_val.font = bold_font
        tot_val.number_format = 'Bs. #,##0.00'
        tot_val.alignment = Alignment(horizontal="right", vertical="center")

        # Auto-ajuste de anchos de columna
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def generate_ventas_pdf(
        self,
        fecha_inicio: Optional[date] = None,
        fecha_fin: Optional[date] = None,
        sucursal_id: Optional[int] = None,
        metodo_pago: Optional[str] = None,
    ) -> io.BytesIO:
        """Genera un archivo PDF con el reporte de ventas."""
        ventas = self.get_ventas_data(fecha_inicio, fecha_fin, sucursal_id, metodo_pago)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#14263D')
        )
        subtitle_style = ParagraphStyle(
            'SubTitleStyle',
            parent=styles['Normal'],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#666666')
        )

        elements: List[Any] = []
        elements.append(Paragraph("StyleStore - Reporte de Ventas", title_style))
        gen_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        elements.append(Paragraph(f"Emitido el: {gen_str} | Ventas listadas: {len(ventas)}", subtitle_style))
        elements.append(Spacer(1, 15))

        # Tabla de ventas
        table_data = [["N° Ticket / Orden", "Fecha", "Cliente / Correo", "Sucursal", "Método", "Total (Bs.)"]]
        gran_total = Decimal("0.00")

        for v in ventas:
            gran_total += Decimal(str(v.total))
            cli_nom = _get_cliente_nombre(v.cliente)
            cli_email = _get_cliente_email(v.cliente)
            suc_nom = _get_sucursal_nombre(v.sucursal)
            cliente_cell = f"{cli_nom}\n{cli_email}" if cli_email else cli_nom

            table_data.append([
                v.ticket_numero or f"ORD-{v.id:04d}",
                v.fecha.strftime("%d/%m/%Y") if v.fecha else "",
                cliente_cell,
                suc_nom[:18],
                (v.metodo_pago or "EFECTIVO")[:10].upper(),
                f"Bs. {float(v.total):.2f}"
            ])

        table_data.append(["", "", "", "", "TOTAL:", f"Bs. {float(gran_total):.2f}"])

        col_widths = [90, 65, 145, 95, 65, 80]
        table = Table(table_data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#14263D')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.HexColor('#F8F9FA'), colors.white]),
            ('LINEBELOW', (0, -1), (-1, -1), 1.5, colors.HexColor('#14263D')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ]))

        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return buffer

    # ==========================================
    # 2. REPORTES DE INVENTARIO Y STOCK CRÍTICO
    # ==========================================

    def get_inventario_data(self, sucursal_id: Optional[int] = None, solo_bajo_stock: bool = False):
        """Obtiene el stock actual agrupado con producto, color, talla y sucursal."""
        query = self.db.query(StockInventario).options(
            joinedload(StockInventario.sucursal),
            joinedload(StockInventario.talla),
            joinedload(StockInventario.producto_color).joinedload(ProductoColor.producto),
            joinedload(StockInventario.producto_color).joinedload(ProductoColor.color),
        )
        if sucursal_id:
            query = query.filter(StockInventario.sucursal_id == sucursal_id)
        if solo_bajo_stock:
            query = query.filter(StockInventario.cantidad <= 5)

        return query.all()

    def generate_inventario_excel(self, sucursal_id: Optional[int] = None, solo_bajo_stock: bool = False) -> io.BytesIO:
        """Genera un archivo Excel con el reporte de inventario."""
        items = self.get_inventario_data(sucursal_id, solo_bajo_stock)

        wb = openpyxl.Workbook()
        ws: Any = wb.active if wb.active is not None else wb.create_sheet()
        assert ws is not None
        ws.title = clean_sheet_title("Inventario StyleStore")

        navy_fill = PatternFill(start_color="14263D", end_color="14263D", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=16, bold=True, color="14263D")
        normal_font = Font(name="Arial", size=10)
        thin_border = Border(
            left=Side(style='thin', color='E0E0E0'),
            right=Side(style='thin', color='E0E0E0'),
            top=Side(style='thin', color='E0E0E0'),
            bottom=Side(style='thin', color='E0E0E0')
        )

        ws.merge_cells("A1:F1")
        ws["A1"] = "StyleStore - Reporte de Inventario y Stock"
        ws["A1"].font = title_font

        ws["A2"] = f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total Registros: {len(items)}"
        ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")

        headers = ["Código", "Producto", "Color", "Talla", "Sucursal", "Stock Actual"]
        ws.row_dimensions[4].height = 24
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=4, column=col_idx, value=h)
            c.fill = navy_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = 5
        for item in items:
            p_color = item.producto_color
            prod = p_color.producto if p_color else None
            color_nom = p_color.color.nombre if (p_color and p_color.color) else "-"
            talla_nom = item.talla.nombre if item.talla else "-"

            ws.cell(row=row_idx, column=1, value=prod.codigo if prod else "-").font = normal_font
            ws.cell(row=row_idx, column=2, value=prod.nombre if prod else "-").font = normal_font
            ws.cell(row=row_idx, column=3, value=color_nom).font = normal_font
            ws.cell(row=row_idx, column=4, value=talla_nom).font = normal_font
            ws.cell(row=row_idx, column=5, value=item.sucursal.nombre if item.sucursal else "Central").font = normal_font
            
            c_stock = ws.cell(row=row_idx, column=6, value=item.cantidad)
            c_stock.font = normal_font
            c_stock.alignment = Alignment(horizontal="right")
            if item.cantidad <= 3:
                c_stock.fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")

            for c in range(1, 7):
                ws.cell(row=row_idx, column=c).border = thin_border
            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def generate_inventario_pdf(self, sucursal_id: Optional[int] = None, solo_bajo_stock: bool = False) -> io.BytesIO:
        """Genera un archivo PDF con el inventario."""
        items = self.get_inventario_data(sucursal_id, solo_bajo_stock)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle('InvTitle', parent=styles['Heading1'], fontSize=18, textColor=colors.HexColor('#14263D'))
        elements: List[Any] = [
            Paragraph("StyleStore - Reporte de Inventario", title_style),
            Spacer(1, 15)
        ]

        table_data = [["Código", "Producto", "Color", "Talla", "Sucursal", "Stock"]]
        for it in items:
            pc = it.producto_color
            prod = pc.producto if pc else None
            table_data.append([
                prod.codigo if prod else "-",
                (prod.nombre if prod else "-")[:25],
                pc.color.nombre[:10] if (pc and pc.color) else "-",
                it.talla.nombre[:6] if it.talla else "-",
                (it.sucursal.nombre if it.sucursal else "Central")[:15],
                str(it.cantidad)
            ])

        col_widths = [70, 180, 70, 50, 100, 60]
        table = Table(table_data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#14263D')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F8F9FA'), colors.white]),
        ]))
        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return buffer

    # ==========================================
    # 3. GENERADORES GENÉRICOS (EXCEL & PDF)
    # ==========================================

    def build_generic_excel(self, sheet_title: str, report_title: str, headers: List[str], rows: List[List[Any]]) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws: Any = wb.active if wb.active is not None else wb.create_sheet()
        assert ws is not None
        ws.title = clean_sheet_title(sheet_title)

        navy_fill = PatternFill(start_color="14263D", end_color="14263D", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=15, bold=True, color="14263D")
        normal_font = Font(name="Arial", size=10)
        thin_border = Border(
            left=Side(style='thin', color='E0E0E0'),
            right=Side(style='thin', color='E0E0E0'),
            top=Side(style='thin', color='E0E0E0'),
            bottom=Side(style='thin', color='E0E0E0')
        )

        num_cols = len(headers)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        ws.cell(row=1, column=1, value=f"StyleStore - {report_title}").font = title_font

        ws.cell(row=2, column=1, value=f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total: {len(rows)}").font = Font(name="Arial", size=9, italic=True, color="666666")

        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=c_idx, value=h)
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        curr_row = 5
        for row in rows:
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=curr_row, column=c_idx, value=val)
                cell.font = normal_font
                cell.border = thin_border
                if isinstance(val, (int, float, Decimal)):
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(horizontal="left")
            curr_row += 1

        for col in ws.columns:
            max_len = max(len(str(c.value or '')) for c in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def build_generic_pdf(self, report_title: str, headers: List[str], rows: List[List[Any]], col_widths: Optional[List[int]] = None) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle('RepTitle', parent=styles['Heading1'], fontSize=16, leading=20, textColor=colors.HexColor('#14263D'))
        sub_style = ParagraphStyle('RepSub', parent=styles['Normal'], fontSize=8, textColor=colors.HexColor('#666666'))

        elements: List[Any] = [
            Paragraph(f"StyleStore - {report_title}", title_style),
            Paragraph(f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Registros: {len(rows)}", sub_style),
            Spacer(1, 12),
        ]

        table_data = [headers]
        for r in rows:
            table_data.append([str(c) if c is not None else "-" for c in r])

        table = Table(table_data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#14263D')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F8F9FA'), colors.white]),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return buffer

    # ==========================================
    # 4. LOS 8 NUEVOS REPORTES ESPECÍFICOS (PUNTO 11)
    # ==========================================

    # Reporte 1: Productos más vendidos
    def get_mas_vendidos_data(self, limit: int = 20) -> List[Dict[str, Any]]:
        q = (
            self.db.query(
                DetalleVenta.producto_nombre,
                func.sum(DetalleVenta.cantidad).label("unidades"),
                func.sum(DetalleVenta.subtotal).label("monto"),
            )
            .group_by(DetalleVenta.producto_nombre)
            .order_by(func.sum(DetalleVenta.cantidad).desc())
            .limit(limit)
            .all()
        )
        return [
            {
                "producto": r[0] or "Sin nombre",
                "unidades_vendidas": int(r[1] or 0),
                "total_recaudado": float(r[2] or 0.0),
            }
            for r in q
        ]

    # Reporte 2: Ventas por producto
    def get_ventas_por_producto_data(self) -> List[Dict[str, Any]]:
        q = (
            self.db.query(
                DetalleVenta.producto_nombre,
                DetalleVenta.color_nombre,
                DetalleVenta.talla_nombre,
                func.sum(DetalleVenta.cantidad).label("unidades"),
                func.sum(DetalleVenta.subtotal).label("monto"),
            )
            .group_by(DetalleVenta.producto_nombre, DetalleVenta.color_nombre, DetalleVenta.talla_nombre)
            .order_by(DetalleVenta.producto_nombre.asc())
            .all()
        )
        return [
            {
                "producto": r[0] or "-",
                "color": r[1] or "-",
                "talla": r[2] or "-",
                "unidades": int(r[3] or 0),
                "subtotal": float(r[4] or 0.0),
            }
            for r in q
        ]

    # Reporte 3: Ventas por tipo (en línea vs presencial)
    def get_ventas_por_tipo_data(self) -> List[Dict[str, Any]]:
        q = (
            self.db.query(
                OrdenVenta.tipo_venta,
                func.count(OrdenVenta.id).label("total_ordenes"),
                func.sum(OrdenVenta.total).label("total_monto"),
            )
            .filter(OrdenVenta.estado != "anulado")
            .group_by(OrdenVenta.tipo_venta)
            .all()
        )
        return [
            {
                "tipo_venta": (r[0] or "presencial").capitalize(),
                "total_ordenes": int(r[1] or 0),
                "total_monto": float(r[2] or 0.0),
            }
            for r in q
        ]

    # Reporte 4: Clientes registrados con total de compras
    def get_clientes_registrados_data(self) -> List[Dict[str, Any]]:
        clientes = self.db.query(Cliente).options(joinedload(Cliente.user)).all()
        resultado = []
        for c in clientes:
            nom = f"{c.user.name} {c.user.apellido or ''}".strip() if (c.user and c.user.name) else (c.nombre or "-")
            email = c.user.email if (c.user and c.user.email) else "-"
            # Compras del cliente
            compras_q = self.db.query(OrdenVenta).filter(OrdenVenta.codigo_cliente == c.codigo, OrdenVenta.estado != "anulado").all()
            total_compras = len(compras_q)
            monto_gastado = sum(float(x.total) for x in compras_q)
            resultado.append({
                "codigo": c.codigo,
                "nombre": nom,
                "email": email,
                "telefono": c.telefono or "-",
                "total_compras": total_compras,
                "monto_gastado": round(monto_gastado, 2),
            })
        resultado.sort(key=lambda x: x["monto_gastado"], reverse=True)
        return resultado

    # Reporte 5: Empleados por sucursal
    def get_empleados_por_sucursal_data(self, sucursal_id: Optional[int] = None) -> List[Dict[str, Any]]:
        q = self.db.query(Empleado).options(joinedload(Empleado.user), joinedload(Empleado.sucursal))
        if sucursal_id:
            q = q.filter(Empleado.sucursal_id == sucursal_id)
        empleados = q.all()
        return [
            {
                "codigo": e.codigo,
                "nombre": f"{e.user.name} {e.user.apellido or ''}".strip() if e.user else "-",
                "email": e.user.email if e.user else "-",
                "sucursal": e.sucursal.nombre if e.sucursal else "-",
                "telefono": e.telefono or "-",
                "sueldo": float(e.sueldo),
            }
            for e in empleados
        ]

    # Reporte 6: Rotación de prendas (Top 5 más y Top 5 menos vendidas)
    def get_rotacion_prendas_data(self) -> Dict[str, Any]:
        mas = self.get_mas_vendidos_data(limit=5)
        # Menos vendidas
        q_menos = (
            self.db.query(
                DetalleVenta.producto_nombre,
                func.sum(DetalleVenta.cantidad).label("unidades"),
                func.sum(DetalleVenta.subtotal).label("monto"),
            )
            .group_by(DetalleVenta.producto_nombre)
            .order_by(func.sum(DetalleVenta.cantidad).asc())
            .limit(5)
            .all()
        )
        menos = [
            {
                "producto": r[0] or "Sin nombre",
                "unidades_vendidas": int(r[1] or 0),
                "total_recaudado": float(r[2] or 0.0),
            }
            for r in q_menos
        ]
        return {"mas_vendidas": mas, "menos_vendidas": menos}

    # Reporte 7: Reporte de envíos (Delivery StyleStore)
    def get_envios_report_data(self) -> List[Dict[str, Any]]:
        envios = self.db.query(Envio).options(
            joinedload(Envio.orden_venta).joinedload(OrdenVenta.cliente).joinedload(Cliente.user)
        ).order_by(Envio.id.desc()).all()
        res = []
        for e in envios:
            cli_nom = "Cliente General"
            if e.orden_venta and e.orden_venta.cliente and e.orden_venta.cliente.user:
                u = e.orden_venta.cliente.user
                cli_nom = f"{u.name} {u.apellido or ''}".strip()
            ticket = e.orden_venta.ticket_numero if e.orden_venta else f"ORD-{e.orden_venta_id}"

            res.append({
                "id": e.id,
                "ticket": ticket,
                "cliente": cli_nom,
                "ciudad": e.ciudad,
                "direccion": e.direccion,
                "conductor": getattr(e, "delivery_conductor", None) or "Por asignar",
                "tracking": getattr(e, "tracking_code", None) or getattr(e, "token_seguimiento", None) or getattr(e, "yango_tracking_code", None) or "-",
                "estado": e.estado.upper(),
                "costo": float(e.costo),
                "fecha": e.fecha.strftime("%d/%m/%Y") if e.fecha else "-",
            })
        return res

    # Reporte 8: Devoluciones y Cambios
    def get_cambios_devoluciones_report_data(self, sucursal_id: Optional[int] = None) -> List[Dict[str, Any]]:
        q = self.db.query(CambioDevolucion).options(
            joinedload(CambioDevolucion.orden_venta),
            joinedload(CambioDevolucion.detalle_venta),
            joinedload(CambioDevolucion.sucursal),
        )
        if sucursal_id:
            q = q.filter(CambioDevolucion.sucursal_id == sucursal_id)

        items = q.order_by(CambioDevolucion.id.desc()).all()
        res = []
        for c in items:
            ticket = c.orden_venta.ticket_numero if c.orden_venta else f"ORD-{c.orden_venta_id}"
            prod_nom = c.detalle_venta.producto_nombre if c.detalle_venta else "-"
            suc_nom = c.sucursal.nombre if c.sucursal else "-"

            res.append({
                "id": c.id,
                "ticket": ticket,
                "tipo": c.tipo.upper(),
                "motivo": c.motivo,
                "producto": prod_nom,
                "sucursal": suc_nom,
                "estado": c.estado.upper(),
                "fecha_programada": c.fecha_programada.strftime("%d/%m/%Y") if c.fecha_programada else "-",
            })
        return res

    # Reporte Financiero: Flujo de caja y pagos por método
    def get_financiero_report_data(self, sucursal_id: Optional[int] = None) -> List[Dict[str, Any]]:
        q = self.db.query(Pago).join(OrdenVenta, Pago.orden_venta_id == OrdenVenta.id)
        if sucursal_id:
            q = q.filter(OrdenVenta.sucursal_id == sucursal_id)
        pagos = q.all()

        agrupado: Dict[str, Dict[str, Any]] = {}
        for p in pagos:
            metodo = (getattr(p, "metodo_pago", None) or getattr(p, "tipo_pago", None) or "EFECTIVO").upper()
            if metodo not in agrupado:
                agrupado[metodo] = {
                    "metodo": metodo,
                    "cantidad_transacciones": 0,
                    "total_recaudado": 0.0,
                }
            agrupado[metodo]["cantidad_transacciones"] += 1
            agrupado[metodo]["total_recaudado"] += float(p.monto)

        return list(agrupado.values())

    # Reporte de Caducidad / Obsolescencia y Temporadas
    def get_caducidad_report_data(self, sucursal_id: Optional[int] = None) -> List[Dict[str, Any]]:
        # Temporadas activas
        temporadas_activas = {t.id for t in self.db.query(Temporada).filter(Temporada.activa == True).all()}
        
        q = (
            self.db.query(StockInventario)
            .join(ProductoColor, StockInventario.producto_color_id == ProductoColor.id)
            .join(Producto, ProductoColor.producto_codigo == Producto.codigo)
            .options(
                joinedload(StockInventario.talla),
                joinedload(StockInventario.sucursal),
            )
            .filter(StockInventario.cantidad > 0)
        )
        if sucursal_id:
            q = q.filter(StockInventario.sucursal_id == sucursal_id)

        items = q.all()
        res = []
        for s in items:
            prod = s.producto_color.producto if s.producto_color else None
            if not prod:
                continue
            # Prenda de temporada pasada o inactiva
            if prod.temporada_id and prod.temporada_id not in temporadas_activas:
                temp_nom = prod.temporada.nombre if prod.temporada else "Pasada"
                suc_nom = s.sucursal.nombre if s.sucursal else "General"
                res.append({
                    "producto": prod.nombre,
                    "codigo": prod.codigo,
                    "temporada": temp_nom,
                    "sucursal": suc_nom,
                    "talla": s.talla.nombre if s.talla else "-",
                    "unidades_stock": s.cantidad,
                    "precio_actual": float(prod.precio),
                    "descuento_sugerido": "30% - 50% (Liquidación)",
                })
        return res

    # Reporte de Auditoría y Bitácora
    def get_auditoria_report_data(self, limit: int = 150) -> List[Dict[str, Any]]:
        logs = (
            self.db.query(Bitacora)
            .options(joinedload(Bitacora.user))
            .order_by(Bitacora.created_at.desc())
            .limit(limit)
            .all()
        )
        res = []
        for b in logs:
            res.append({
                "id": b.id,
                "usuario": f"{b.user.name} ({b.user.role})" if b.user else "Sistema",
                "modulo": b.module.upper() if b.module else "SISTEMA",
                "accion": b.action,
                "ip": getattr(b, "ip_address", None) or "127.0.0.1",
                "fecha": b.created_at.strftime("%d/%m/%Y %H:%M:%S") if b.created_at else "-",
            })
        return res

    # Reporte de Compras a Proveedores
    def get_compras_report_data(self, sucursal_id: Optional[int] = None) -> List[Dict[str, Any]]:
        try:
            from app.models.compra import Compra
            q = self.db.query(Compra).options(joinedload(Compra.proveedor), joinedload(Compra.sucursal))
            if sucursal_id:
                q = q.filter(Compra.sucursal_id == sucursal_id)
            compras = q.order_by(Compra.fecha.desc()).all()
            res = []
            for c in compras:
                res.append({
                    "id": c.id,
                    "fecha": c.fecha.strftime("%d/%m/%Y") if c.fecha else "-",
                    "proveedor": c.proveedor.nombre if c.proveedor else "Proveedor General",
                    "sucursal": c.sucursal.nombre if c.sucursal else "General",
                    "total": float(c.total),
                })
            return res
        except Exception:
            return []

